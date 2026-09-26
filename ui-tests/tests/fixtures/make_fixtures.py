"""Regenerate the documents the integration tests open.

The tests assert on the text written here, so change both together.
Needs python-docx, python-pptx and openpyxl: `python make_fixtures.py` in
this folder. ODT and ODP are written as plain zip archives.
"""
import zipfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation

HERE = Path(__file__).parent


def add_link(paragraph, text, target):
    """Append a hyperlink run; python-docx has no public API for it"""
    rel = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    label = OxmlElement("w:t")
    label.text = text
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def make_docx():
    doc = Document()
    doc.add_heading("Quarterly Report", level=1)
    para = doc.add_paragraph("Revenue grew in ")
    para.add_run("every region").bold = True
    table = doc.add_table(rows=2, cols=2)
    for row, (region, revenue) in zip(table.rows, [("Region", "Revenue"), ("North", "120")]):
        row.cells[0].text, row.cells[1].text = region, revenue
    links = doc.add_paragraph()
    add_link(links, "Jupyter website", "https://jupyter.org")
    links.add_run(" and ")
    add_link(links, "Unsafe link", "javascript:alert(1)")
    doc.add_page_break()
    doc.add_heading("Appendix", level=1)
    doc.add_paragraph("Second page text.")
    doc.save(HERE / "sample.docx")


def make_altchunk_docx():
    """A DOCX embedding an HTML part whose script marks the page if it runs"""
    doc = Document()
    doc.add_paragraph("Text around the chunk")
    html = (
        "<html><body><p>Embedded chunk text</p>"
        "<script>parent.document.body.dataset.altchunk = 'ran'</script>"
        "</body></html>"
    ).encode()
    part = Part(PackURI("/word/chunk.html"), "text/html", html, doc.part.package)
    chunk = OxmlElement("w:altChunk")
    chunk.set(qn("r:id"), doc.part.relate_to(part, RELATIONSHIP_TYPE.A_F_CHUNK))
    doc.paragraphs[-1]._p.addnext(chunk)
    doc.save(HERE / "altchunk.docx")


def make_pptx():
    deck = Presentation()
    texts = [
        ("Alpha slide", "First slide"),
        ("Beta slide", "Second slide"),
        ("Gamma slide", "The zebra appears only here"),
    ]
    for title, subtitle in texts:
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = subtitle
    deck.save(HERE / "sample.pptx")


ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
    'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
    'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
    'xmlns:xlink="http://www.w3.org/1999/xlink" office:version="1.2"'
)


def write_odf(name, mime, styles, automatic, body):
    """Write an ODF package: the uncompressed mimetype entry comes first"""
    content = (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ODF_NS}>'
        f"<office:automatic-styles>{automatic}</office:automatic-styles>"
        f"<office:body>{body}</office:body></office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
        'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
        f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    with zipfile.ZipFile(HERE / name, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), mime, zipfile.ZIP_STORED)
        archive.writestr("content.xml", content)
        archive.writestr("styles.xml", f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {ODF_NS}>{styles}</office:document-styles>')
        archive.writestr("META-INF/manifest.xml", manifest)


def make_odt():
    styles = (
        '<office:styles><style:style style:name="Heading_20_1" style:family="paragraph">'
        '<style:text-properties fo:font-size="20pt" fo:font-weight="bold"/></style:style></office:styles>'
        '<office:automatic-styles><style:page-layout style:name="pm1"><style:page-layout-properties '
        'fo:page-width="21cm" fo:page-height="29.7cm" fo:margin="2cm"/></style:page-layout></office:automatic-styles>'
        '<office:master-styles><style:master-page style:name="Standard" style:page-layout-name="pm1"/></office:master-styles>'
    )
    automatic = (
        '<style:style style:name="T1" style:family="text"><style:text-properties fo:font-weight="bold"/></style:style>'
        '<style:style style:name="P1" style:family="paragraph"><style:paragraph-properties fo:break-before="page"/></style:style>'
    )
    cells = "".join(
        "<table:table-row>"
        + "".join(f"<table:table-cell><text:p>Cell {column}{row}</text:p></table:table-cell>" for column in "AB")
        + "</table:table-row>"
        for row in "12"
    )
    body = (
        "<office:text>"
        '<text:h text:style-name="Heading_20_1" text:outline-level="1">Sample Heading</text:h>'
        '<text:p>Plain text before a <text:span text:style-name="T1">bold run</text:span> and after.</text:p>'
        f'<table:table table:name="Table1"><table:table-column table:number-columns-repeated="2"/>{cells}</table:table>'
        '<text:p><text:a xlink:type="simple" xlink:href="https://example.org/">Example link</text:a> and '
        '<text:a xlink:type="simple" xlink:href="javascript:parent.document.body.dataset.odflink=1">Unsafe link</text:a></text:p>'
        '<text:p text:style-name="P1">Second page text.</text:p>'
        "</office:text>"
    )
    write_odf("sample.odt", "application/vnd.oasis.opendocument.text", styles, automatic, body)


def make_odp():
    styles = (
        '<office:automatic-styles><style:page-layout style:name="PM1"><style:page-layout-properties '
        'fo:page-width="28cm" fo:page-height="15.75cm"/></style:page-layout></office:automatic-styles>'
        '<office:master-styles><style:master-page style:name="Default" style:page-layout-name="PM1"/></office:master-styles>'
    )
    automatic = '<style:style style:name="P1" style:family="paragraph"><style:text-properties fo:font-size="36pt" fo:font-weight="bold"/></style:style>'
    slides = "".join(
        f'<draw:page draw:name="page{n}" draw:master-page-name="Default">'
        '<draw:frame svg:width="24cm" svg:height="3cm" svg:x="2cm" svg:y="1cm" presentation:class="title">'
        f'<draw:text-box><text:p text:style-name="P1">{title}</text:p></draw:text-box></draw:frame>'
        '<draw:frame svg:width="24cm" svg:height="9cm" svg:x="2cm" svg:y="5cm" presentation:class="outline">'
        f"<draw:text-box><text:p>{body}</text:p></draw:text-box></draw:frame></draw:page>"
        for n, (title, body) in enumerate(
            [("Alpha", "First slide"), ("Beta", "Second slide"), ("Gamma", "The zebra appears only here")], 1
        )
    )
    write_odf("sample.odp", "application/vnd.oasis.opendocument.presentation", styles, automatic,
              f"<office:presentation>{slides}</office:presentation>")


def make_xlsx():
    """Values only: the viewer shows saved formula results, and openpyxl saves none"""
    book = Workbook()
    data = book.active
    data.title = "Data"
    for row in [("Item", "Quantity", "Price"), ("Apples", 3, 1.5), ("Pears", 5, 2.25)]:
        data.append(row)
    for cell in data[1]:
        cell.font = Font(bold=True)
    book.create_sheet("Other")["A1"] = "Second sheet cell"
    book.save(HERE / "sample.xlsx")


def make_others():
    (HERE / "sample.rtf").write_text(
        r"{\rtf1\ansi\deff0{\fonttbl{\f0 Arial;}}\f0\fs24 Plain text and {\b bold text} in RTF.\par}"
    )
    # OLE compound-file signature: what a real .doc or .ppt starts with
    ole = bytes.fromhex("d0cf11e0a1b11ae1") + bytes(504)
    (HERE / "legacy.doc").write_bytes(ole)
    (HERE / "legacy.ppt").write_bytes(ole)
    (HERE / "broken.docx").write_bytes(b"This is not a DOCX file.")
    (HERE / "broken.odt").write_bytes(b"This is not an ODT file.")
    # a truncated workbook: it still starts with the zip signature
    workbook = (HERE / "sample.xlsx").read_bytes()
    (HERE / "broken.xlsx").write_bytes(workbook[: len(workbook) // 2])
    (HERE / "empty.pptx").write_bytes(b"")


if __name__ == "__main__":
    make_docx()
    make_altchunk_docx()
    make_pptx()
    make_odt()
    make_odp()
    make_xlsx()
    make_others()
